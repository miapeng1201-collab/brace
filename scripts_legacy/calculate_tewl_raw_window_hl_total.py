from __future__ import annotations

import statistics
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from process_tewl_data import (  # noqa: E402
    ROOT as PROCESS_ROOT,
    NS,
    SOURCE_DIR,
    col_number,
    read_shared_strings,
    read_sheet,
    sheet_paths,
)

OUTPUT_DIR = ROOT / "processed_data" / "04_tewameter"
INPUT_SUMMARY = OUTPUT_DIR / "TEWL_processed_raw_window.xlsx"
OUTPUT_FILE = OUTPUT_DIR / "TEWL_raw_window_HL_total_35_55s.xlsx"

WINDOW_START = 35.0
WINDOW_END = 55.0

SUMMARY_COLUMNS = [
    "subject",
    "group",
    "stage",
    "day",
    "visit",
    "source_file",
    "raw_window_start_s",
    "raw_window_end_s",
    "raw_window_measurement_count",
    "raw_window_total_points",
    "raw_window_hl_hd_mean_of_measurements_w_m2",
    "raw_window_hl_ec_mean_of_measurements_w_m2",
    "raw_window_hl_total_mean_of_measurements_w_m2",
    "hl_total_measurement_1",
    "hl_total_measurement_2",
    "hl_total_measurement_3",
    "hl_total_measurement_4",
    "hl_total_measurement_5",
    "hl_total_measurement_6",
    "hl_total_measurement_7",
    "hl_total_measurement_8",
    "hl_total_measurement_9",
    "hl_total_measurement_10",
]

MEASUREMENT_COLUMNS = [
    "subject",
    "group",
    "stage",
    "day",
    "visit",
    "measurement_number",
    "measurement_label",
    "source_file",
    "window_start_s",
    "window_end_s",
    "window_n_points",
    "window_hl_hd_mean_w_m2",
    "window_hl_ec_mean_w_m2",
    "window_hl_total_mean_w_m2",
]


def number_or_blank(value: Any) -> float | None:
    if value in ("", None):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def read_clean_metadata() -> dict[str, dict[str, Any]]:
    wb = openpyxl.load_workbook(INPUT_SUMMARY, read_only=True, data_only=True)
    ws = wb["RawWindowSummaryClean"]
    rows = ws.iter_rows(values_only=True)
    headers = [str(v).strip() if v is not None else "" for v in next(rows)]
    by_source: dict[str, dict[str, Any]] = {}
    for row in rows:
        item = {headers[i]: row[i] if i < len(row) else "" for i in range(len(headers))}
        source_file = item.get("source_file")
        if source_file:
            by_source[str(source_file)] = item
    wb.close()
    return by_source


def parse_hl_raw_window(path: Path, meta: dict[str, Any]) -> list[dict[str, Any]]:
    rel = str(path.relative_to(PROCESS_ROOT))
    with path.open("rb") as fh:
        import zipfile

        with zipfile.ZipFile(fh) as zf:
            shared_strings = read_shared_strings(zf)
            paths = sheet_paths(zf)
            raw_rows = read_sheet(zf, paths["Raw Data"], shared_strings)

    if len(raw_rows) < 3:
        return []

    label_row = raw_rows[0]
    header_row = raw_rows[1]
    starts: list[tuple[int, str]] = []
    for index, value in enumerate(label_row):
        if isinstance(value, str) and value.startswith("Measurement_"):
            starts.append((index, value))

    results: list[dict[str, Any]] = []
    for pos, (start_col, label) in enumerate(starts):
        end_col = starts[pos + 1][0] if pos + 1 < len(starts) else max(len(row) for row in raw_rows)
        headers = header_row[start_col:end_col]
        try:
            elapsed_idx = start_col + headers.index("Elapsed Time [s]")
            hd_idx = start_col + headers.index("HL by Heat Diffusion [W/m²]")
            ec_idx = start_col + headers.index("HL by Evaporation Cooling [W/m²]")
        except ValueError:
            continue

        hd_values: list[float] = []
        ec_values: list[float] = []
        total_values: list[float] = []
        for row in raw_rows[2:]:
            elapsed = number_or_blank(row[elapsed_idx] if elapsed_idx < len(row) else "")
            hd = number_or_blank(row[hd_idx] if hd_idx < len(row) else "")
            ec = number_or_blank(row[ec_idx] if ec_idx < len(row) else "")
            if elapsed is None or hd is None or ec is None:
                continue
            if WINDOW_START <= elapsed <= WINDOW_END:
                hd_values.append(hd)
                ec_values.append(ec)
                total_values.append(hd + ec)

        try:
            measurement_number = int(label.rsplit("_", 1)[1])
        except (IndexError, ValueError):
            measurement_number = ""

        results.append(
            {
                "subject": meta.get("subject", ""),
                "group": meta.get("group", ""),
                "stage": meta.get("stage", ""),
                "day": meta.get("day", ""),
                "visit": meta.get("visit", ""),
                "measurement_number": measurement_number,
                "measurement_label": label,
                "source_file": rel,
                "window_start_s": WINDOW_START,
                "window_end_s": WINDOW_END,
                "window_n_points": len(total_values),
                "window_hl_hd_mean_w_m2": statistics.fmean(hd_values) if hd_values else "",
                "window_hl_ec_mean_w_m2": statistics.fmean(ec_values) if ec_values else "",
                "window_hl_total_mean_w_m2": statistics.fmean(total_values) if total_values else "",
            }
        )
    return results


def make_summary(measurement_rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    by_source: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in measurement_rows:
        by_source[row["source_file"]].append(row)

    summaries: list[dict[str, Any]] = []
    for source_file, rows in sorted(by_source.items()):
        rows = sorted(rows, key=lambda r: r["measurement_number"] if isinstance(r["measurement_number"], int) else 999)
        first = rows[0]
        hd_means = [float(r["window_hl_hd_mean_w_m2"]) for r in rows if r["window_hl_hd_mean_w_m2"] != ""]
        ec_means = [float(r["window_hl_ec_mean_w_m2"]) for r in rows if r["window_hl_ec_mean_w_m2"] != ""]
        total_means = [float(r["window_hl_total_mean_w_m2"]) for r in rows if r["window_hl_total_mean_w_m2"] != ""]

        measurement_cols = {f"hl_total_measurement_{i}": "" for i in range(1, 11)}
        for row in rows:
            number = row.get("measurement_number")
            if isinstance(number, int) and 1 <= number <= 10:
                measurement_cols[f"hl_total_measurement_{number}"] = row["window_hl_total_mean_w_m2"]

        summaries.append(
            {
                "subject": first["subject"],
                "group": first["group"],
                "stage": first["stage"],
                "day": first["day"],
                "visit": first["visit"],
                "source_file": source_file,
                "raw_window_start_s": WINDOW_START,
                "raw_window_end_s": WINDOW_END,
                "raw_window_measurement_count": len(rows),
                "raw_window_total_points": sum(int(r["window_n_points"]) for r in rows),
                "raw_window_hl_hd_mean_of_measurements_w_m2": statistics.fmean(hd_means) if hd_means else "",
                "raw_window_hl_ec_mean_of_measurements_w_m2": statistics.fmean(ec_means) if ec_means else "",
                "raw_window_hl_total_mean_of_measurements_w_m2": statistics.fmean(total_means) if total_means else "",
                **measurement_cols,
            }
        )
    return summaries


def write_sheet(ws: openpyxl.worksheet.worksheet.Worksheet, rows: list[dict[str, Any]], columns: list[str]) -> None:
    ws.append(columns)
    for row in rows:
        ws.append([row.get(col, "") for col in columns])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main() -> None:
    metadata = read_clean_metadata()
    measurement_rows: list[dict[str, Any]] = []
    errors: list[tuple[str, str]] = []

    for source_file, meta in sorted(metadata.items()):
        path = ROOT / source_file
        if not path.exists():
            errors.append((source_file, "file not found"))
            continue
        try:
            measurement_rows.extend(parse_hl_raw_window(path, meta))
        except Exception as exc:  # keep going across hundreds of raw files
            errors.append((source_file, str(exc)))

    summary_rows = make_summary(measurement_rows)
    summary_rows.sort(key=lambda r: (str(r["group"]), str(r["subject"]), str(r["day"]), str(r["stage"])))
    measurement_rows.sort(
        key=lambda r: (
            str(r["group"]),
            str(r["subject"]),
            str(r["day"]),
            str(r["stage"]),
            int(r["measurement_number"]) if isinstance(r["measurement_number"], int) else 999,
        )
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "HLTotalSummary35_55s"
    write_sheet(ws, summary_rows, SUMMARY_COLUMNS)
    ws2 = wb.create_sheet("HLTotalMeasurements35_55s")
    write_sheet(ws2, measurement_rows, MEASUREMENT_COLUMNS)
    ws3 = wb.create_sheet("QC")
    ws3.append(["item", "value"])
    ws3.append(["source_clean_rows", len(metadata)])
    ws3.append(["summary_rows", len(summary_rows)])
    ws3.append(["measurement_rows", len(measurement_rows)])
    ws3.append(["errors", len(errors)])
    ws3.append(["window_rule", "Elapsed Time [s] between 35 and 55 inclusive"])
    ws3.append(["hl_total_rule", "HL by Heat Diffusion [W/m²] + HL by Evaporation Cooling [W/m²] at each raw point"])
    for source_file, error in errors:
        ws3.append([source_file, error])

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT_FILE)
    print(OUTPUT_FILE)
    print(f"summary_rows={len(summary_rows)} measurement_rows={len(measurement_rows)} errors={len(errors)}")


if __name__ == "__main__":
    main()
